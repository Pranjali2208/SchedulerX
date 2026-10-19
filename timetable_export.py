"""PDF / Word export of the generated timetable.

Both exports are built from one shared model (``build_model``) so that they always
contain the same entries in the same order.  The layout follows the institute's
"Draft Time Table" format: A4 portrait, page border, logo header, title block,
Day & Date / Time / Branch / Sem / Course Name / Code table and a signature footer.

Entries are kept strictly separated by programme and sorted by year of study inside
each programme.  Document order follows the institute's reference timetable:
B.Tech, BCA, MCA, BBA, MBA, M.Tech (each First -> Second -> Third -> Final Year).  Nothing is dropped,
merged or duplicated; rows whose programme or year cannot be read reliably are
listed in a clearly labelled "not identified" section instead of being guessed.
"""
import os
import re
from datetime import date, datetime
from io import BytesIO

LOGO_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "rit-logo.png")

INSTITUTE_NAME = "Rajarambapu Institute of Technology"
PREPARED_BY = "M.M.Dige"
CHECKED_BY = "Mrs. M.M.Patil"

# Column widths in points (sum = 537.3 = A4 width minus the reference margins).
COL_WIDTHS = [64.0, 52.0, 158.0, 28.0, 170.0, 65.3]
HEADERS = ["Day & Date", "Time", "Branch", "Sem", "Course Name", "Code"]

PROGRAMMES = [  # (key, heading) in the mandatory order
    ("BTECH", "B.Tech."),
    ("BCA", "BCA"),
    ("MCA", "MCA"),
    ("BBA", "BBA"),
    ("MBA", "MBA"),
    ("MTECH", "M.Tech."),
]
PROGRAMME_RANK = {key: i for i, (key, _) in enumerate(PROGRAMMES)}
PROGRAMME_LABEL = dict(PROGRAMMES)
UNCLASSIFIED = "UNCLASSIFIED"
YEAR_RANK = {"first": 1, "second": 2, "third": 3, "fourth": 4, "final": 4, "fifth": 5}
YEAR_NONE = 99

YEAR_NAME = {
    1: "F.Y.",
    2: "S.Y.",
    3: "T.Y.",
    4: "Final",
    5: "Fifth"
}


def _year_order(programme, year_rank):
    # Custom academic order for every programme: First -> Second -> Third -> Final.
    return year_rank

_CLASS_RE = re.compile(
    r"^\s*(first|second|third|fourth|fifth|final|1st|2nd|3rd|4th|5th|i|ii|iii|iv|v|[1-5])"
    r"\s+(?:year|yr)\.?\s+(.+?)\s*$", re.I)
_YEAR_WORDS = {"1st": "first", "i": "first", "1": "first", "2nd": "second", "ii": "second", "2": "second",
               "3rd": "third", "iii": "third", "3": "third", "4th": "fourth", "iv": "fourth", "4": "fourth",
               "5th": "fifth", "v": "fifth", "5": "fifth"}

BRANCH_ABBREVIATIONS = {
    "computer science and engineering": "CSE",
    "computer science and engineering (artificial intelligence and machine learning)": "CSE (AIML)",
    "computer science and information technology": "CSIT",
    "computer engineering": "Computer",
    "electronics and tele-communication engineering": "E&TC",
    "mechatronics engineering": "Mechatronics",
    "mechanical engineering": "Mechanical",
    "mechanical engineering (thermal engineering)": "Mechanical (Thermal)",
    "robotics and automation": "R&A",
    "civil engineering": "Civil",
    "electrical engineering": "Electrical",
    "automotive technology": "Automotive",
    "structural engineering": "Structural",
    "construction management": "Construction Mgmt",
    "bachelor of computer application": "BCA",
    "bachelor of business administration": "BBA",
    "master of business administration": "MBA",
    "master of computer application": "MCA",
    "m.b.a. innovation, entrepreneurship and venture development": "MBA (IEVD)",
}
ALL_BRANCHES_THRESHOLD = 4  # this many distinct branches or more -> "All" (as in the reference)


# --------------------------------------------------------------------------- model
def _split_outside_parentheses(text):
    parts, depth, current = [], 0, []
    for ch in text:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if ch == "," and depth == 0:
            parts.append("".join(current).strip())
            current = []
        else:
            current.append(ch)
    parts.append("".join(current).strip())
    return [p for p in parts if p]


def _abbreviate(branch):
    cleaned = branch.strip()
    key = cleaned.casefold()
    if key in BRANCH_ABBREVIATIONS:
        return BRANCH_ABBREVIATIONS[key]
    if key.startswith("(") and key.endswith(")") and key[1:-1] in BRANCH_ABBREVIATIONS:
        return BRANCH_ABBREVIATIONS[key[1:-1]]
    return cleaned  # unknown names are shown in full rather than invented


def _branch_text(raw):
    names = []
    for name in _split_outside_parentheses(raw or ""):
        short = _abbreviate(name)
        if short not in names:
            names.append(short)
    if len(names) >= ALL_BRANCHES_THRESHOLD:
        return "All"
    return ", ".join(names)


_YEAR_ONLY_RE = re.compile(
    r"^\s*(first|second|third|fourth|fifth|final|1st|2nd|3rd|4th|5th|i|ii|iii|iv|v|[1-5])"
    r"\s+(?:year|yr)\.?\s*$", re.I)
_ENGINEERING_RE = re.compile(
    r"engineering|technology|robotics|mechatronics|civil|electrical|electronics|computer|"
    r"\bcse\b|\bcsit\b|\be&tc\b|\bentc\b|\betc\b|\baiml\b|\br&a\b|automotive|mechanical|\ball\b", re.I)
_MTECH_ONLY_RE = re.compile(r"structural|thermal|construction", re.I)


def _branch_programme(name):
    """Programme implied by one Branch name, or None when it cannot be told reliably."""
    text = re.sub(r"\s+", " ", name or "").strip().casefold()
    plain = re.sub(r"[^a-z]", "", text)
    if plain.startswith("mba") or "masterofbusiness" in plain:
        return "MBA"
    if plain.startswith("bba") or "bachelorofbusiness" in plain:
        return "BBA"
    if plain.startswith("bca") or "bachelorofcomputerapplication" in plain:
        return "BCA"
    if plain.startswith("mca") or "masterofcomputerapplication" in plain:
        return "MCA"
    if plain.startswith("mtech") or _MTECH_ONLY_RE.search(text):
        return "MTECH"
    if _ENGINEERING_RE.search(text):
        return "BTECH_OR_MTECH"  # an engineering branch name; year/'all' decide which
    return None


def _classify(class_text, branch_text=""):
    """Return (programme_key, year_rank, year_label) from the Class (and Branch) columns, never guessing.

    ``Class`` may carry the programme ("Third Year B.Tech.") or be year-only ("Third Year");
    in the year-only case the programme is read from the Branch column.
    """
    raw = (class_text or "").strip()
    parts = _split_outside_parentheses(raw)
    parsed = []
    year_only = True
    for part in parts:
        match = _CLASS_RE.match(part)
        if match:
            year_only = False
            year_word, programme = match.groups()
            year_word = _YEAR_WORDS.get(year_word.lower(), year_word.lower())
            key = re.sub(r"[^A-Za-z]", "", programme).upper()
            if key not in PROGRAMME_RANK:
                return UNCLASSIFIED, YEAR_NONE, raw
        else:
            match = _YEAR_ONLY_RE.match(part)
            if not match:
                return UNCLASSIFIED, YEAR_NONE, raw
            year_word = _YEAR_WORDS.get(match.group(1).lower(), match.group(1).lower())
            key = None
        parsed.append((key, YEAR_RANK[year_word.lower()], year_word.capitalize()))
    if not parsed:
        return UNCLASSIFIED, YEAR_NONE, raw
    if any(p[0] is None for p in parsed):
        if not year_only:
            return UNCLASSIFIED, YEAR_NONE, raw  # half explicit, half year-only: do not guess
        names = _split_outside_parentheses(branch_text or "")
        kinds = {_branch_programme(n) for n in names}
        if not names or None in kinds:
            return UNCLASSIFIED, YEAR_NONE, raw
        if kinds == {"BTECH_OR_MTECH"} or kinds <= {"BTECH_OR_MTECH", "MTECH"}:
            only_first = all(p[1] == 1 for p in parsed)
            has_all = any(re.sub(r"[^a-z]", "", n.casefold()) == "all" for n in names)
            # First-year B.Tech papers are common to every branch ("all"); a first-year paper
            # with a named branch belongs to M.Tech.  Later years default to B.Tech.
            programme = "MTECH" if ("MTECH" in kinds or (only_first and not has_all)) else "BTECH"
            if "MTECH" in kinds and kinds != {"MTECH"} and not only_first:
                return UNCLASSIFIED, YEAR_NONE, raw
        elif len(kinds) == 1:
            programme = next(iter(kinds))
        else:
            return UNCLASSIFIED, YEAR_NONE, raw  # would mix programmes
        parsed = [(programme, rank, label) for _, rank, label in parsed]
    if len({p[0] for p in parsed}) > 1:
        return UNCLASSIFIED, YEAR_NONE, raw  # would mix programmes: keep apart, do not guess
    parsed.sort(key=lambda p: _year_order(p[0], p[1]))
    years = []
    for _, _, label in parsed:
        if label not in years:
            years.append(label)

        year_abbreviations = {
            1: "F.Y.",
            2: "S.Y.",
            3: "T.Y.",
            4: "Final Year",
            5: "Fifth Year"
        }

        display_years = []
        for _, rank, _ in parsed:
            label = year_abbreviations.get(rank, "Year")
            if label not in display_years:
                display_years.append(label)

        return parsed[0][0], parsed[0][1], " & ".join(display_years)



def _parse_day_date(value):
    text = (value or "").strip()
    day, sep, day_date = text.partition(",")
    if not sep:
        return text, None, text
    day_date = day_date.strip()
    try:
        parsed = datetime.strptime(day_date, "%d.%m.%Y").date()
    except ValueError:
        parsed = None
    return day.strip() + "\n" + day_date, parsed, text


def _time_minutes(value):
    match = re.match(r"\s*(\d{1,2})[.:](\d{2})\s*([ap]m)", value or "", re.I)
    if not match:
        return 0
    hour, minute, meridian = int(match[1]) % 12, int(match[2]), match[3].lower()
    return (hour + (12 if meridian == "pm" else 0)) * 60 + minute


def _time_display(value):
    parts = re.split(r"\s+to\s+", (value or "").strip(), maxsplit=1, flags=re.I)
    return parts[0] + "\nTo\n" + parts[1] if len(parts) == 2 else (value or "")


SUMMARY_PROGRAMME = {"BTECH": "B.Tech.", "BCA": "B.C.A.", "MCA": "M.C.A.", "BBA": "B.B.A.",
                     "MBA": "M.B.A.", "MTECH": "M.Tech."}
SUMMARY_YEAR = {1: "F.Y.", 2: "S.Y.", 3: "T.Y.", 4: "Final Yr.", 5: "Fifth Yr."}
_IEV_RE = re.compile(r"\biev\b|ievd|innovation,?\s+entrepreneurship", re.I)
SUMMARY_MAX_SIZE, SUMMARY_MIN_SIZE, SUMMARY_WIDTH = 7.6, 5.0, 530.0


def build_summary(groups):
    """groups: [(programme_key, year_rank, has_iev)] in document order -> e.g.
    'F.Y.B.Tech., S.Y.B.Tech., T.Y.B.Tech. and Final Yr. B.Tech.' (only what is really present)."""
    items = []
    for key, rank, has_iev in groups:
        year = SUMMARY_YEAR.get(rank)
        name = SUMMARY_PROGRAMME.get(key)
        if not year or not name:
            continue
        label = year + (" " if year.endswith("Yr.") else "") + name
        if has_iev and key == "MBA":
            label += " & IEV"
        if label not in items:
            items.append(label)
    if len(items) > 1:
        return ", ".join(items[:-1]) + " and " + items[-1]
    return items[0] if items else ""


def _summary_width(text, size):
    from reportlab.pdfbase.pdfmetrics import stringWidth
    return stringWidth(text, "Times-Bold", size)


def summary_layout(text):
    """(lines, font size): the summary on one header line, or on two balanced lines when it cannot fit
    one at 7.6 pt (many programmes/years).  Shared by the PDF and Word exports."""
    if not text:
        return [], SUMMARY_MAX_SIZE
    try:
        wide = _summary_width(text, SUMMARY_MAX_SIZE) > SUMMARY_WIDTH
    except ImportError:
        return [text], SUMMARY_MAX_SIZE
    lines = [text]
    if wide:
        cuts = [m.end() for m in re.finditer(r", ", text)]
        best = min(cuts, key=lambda c: abs(_summary_width(text[:c], SUMMARY_MAX_SIZE)
                                           - _summary_width(text[c:], SUMMARY_MAX_SIZE)))
        lines = [text[:best].rstrip(), text[best:]]
    size = SUMMARY_MAX_SIZE
    while size > SUMMARY_MIN_SIZE and max(_summary_width(l, size) for l in lines) > SUMMARY_WIDTH:
        size = round(size - 0.1, 1)
    return lines, size


DEFAULT_ACADEMIC_YEAR = "2026-27"
DEFAULT_SEMESTER = "Sem I"


def _format_semester(value):
    """Normalise 'Sem - 1' / 'Sem - 2' / 'Sem I & Sem II' to 'Sem I' / 'Sem II' style."""
    text = (value or "").strip()
    if not text:
        return DEFAULT_SEMESTER
    roman = {"1": "I", "2": "II", "i": "I", "ii": "II"}
    def one(m):
        return "Sem " + roman.get(m.group(1).lower(), m.group(1).upper())
    return re.sub(r"(?i)\bsem(?:ester)?\s*[-.:]?\s*(1|2|ii|i)\b", one, text)


def _exam_title(exam_type):
    text = (exam_type or "").strip().lower()
    if any(k in text for k in ("re-exam", "re exam", "reexam")):
        return "Re-Examination (Re-exam)"
    return "End Semester Examination (ESE)"


def _header_info(dates, academic_year=None, semester=None, summary="", timetable_status="draft", exam_type=None):
    today = date.today()
    lines, size = summary_layout(summary)
    return {"summary_lines": lines, "summary_size": size,
            "academic_year": (academic_year or "").strip() or DEFAULT_ACADEMIC_YEAR,
            "semester": _format_semester(semester),
            "prepared_date": today.strftime("%d/%m/%Y"),
            "exam_title": _exam_title(exam_type),
            "timetable_status": "Final Time Table" if str(timetable_status).lower() == "final" else "Draft Time Table"}


def build_model(rows, academic_year=None, semester=None, timetable_status="draft", exam_type=None):
    """rows: header row + 8 columns (Day & Date, Time, Branch, Class, Sem, Course Name, Code, Count)."""
    entries, dates = [], []
    for index, row in enumerate(rows[1:]):
        row = list(row) + [""] * (8 - len(row))
        day_date, parsed_date, raw_day = _parse_day_date(row[0])
        if parsed_date:
            dates.append(parsed_date)
        programme, year_rank, year_label = _classify(row[3], row[2])
        branch = _branch_text(row[2])
        if programme == UNCLASSIFIED:
            cell_branch = ((row[3] or "").strip() + "\n" + branch).strip()
        else:
            cell_branch = year_label + " " + branch
        entries.append({
            "index": index,
            "programme": programme,
            "year_rank": year_rank,
            "year_label": year_label if programme != UNCLASSIFIED else "Year not identified",
            "date_key": (parsed_date or date.max, raw_day),
            "time_key": _time_minutes(row[1]),
            "day_date": day_date,
            "time": _time_display(row[1]),
            "block_key": (raw_day, (row[1] or "").strip()),
            "branch": cell_branch,
            "sem": (row[4] or "").strip(),
            "course": (row[5] or "").strip(),
            "code": (row[6] or "").strip(),
            "iev": bool(_IEV_RE.search(row[2] or "")),
        })
    entries.sort(key=lambda e: (
        PROGRAMME_RANK.get(e["programme"], len(PROGRAMMES)), _year_order(e["programme"], e["year_rank"]),
        e["date_key"], e["time_key"], e["index"]))

    summary_groups = []
    for entry in entries:
        if entry["programme"] in SUMMARY_PROGRAMME and entry["year_rank"] in SUMMARY_YEAR:
            group = (entry["programme"], entry["year_rank"])
            if not summary_groups or summary_groups[-1][:2] != group:
                summary_groups.append([*group, False])
            summary_groups[-1][2] = summary_groups[-1][2] or entry["iev"]

    sections = []
    for entry in entries:
        if not sections or sections[-1]["key"] != entry["programme"]:
            heading = PROGRAMME_LABEL.get(entry["programme"], "Programme not identified")
            sections.append({"key": entry["programme"], "heading": heading, "years": []})
        years = sections[-1]["years"]
        if not years or years[-1]["rank"] != entry["year_rank"]:
            years.append({"rank": entry["year_rank"], "blocks": [],
                          
        "label": (
            YEAR_NAME[entry["year_rank"]]
            if entry["year_rank"] in (1, 2, 3)
            else YEAR_NAME[entry["year_rank"]] + " Year"
        ) if entry["year_rank"] in YEAR_NAME else entry["year_label"]})

        blocks = years[-1]["blocks"]
        if not blocks or blocks[-1]["key"] != entry["block_key"]:
            blocks.append({"key": entry["block_key"], "day_date": entry["day_date"],
                           "time": entry["time"], "rows": []})
        blocks[-1]["rows"].append((entry["branch"], entry["sem"], entry["course"], entry["code"]))
    # Year headings: "<Programme> - <Year>" (unclassified rows already carry their own label).
    for section in sections:
        for year in section["years"]:
            if section["key"] == UNCLASSIFIED:
                year["title"] = "Programme / year not identified"
            elif year["rank"] == YEAR_NONE:
                year["title"] = f"{section['heading']} - Year not specified"
            else:
                year["title"] = f"{section['heading']} - {year['label']}"
    return {"header": _header_info(dates, academic_year, semester, build_summary(summary_groups), timetable_status, exam_type), "sections": sections, "total": len(entries)}


def count_entries(model):
    return sum(len(b["rows"]) for s in model["sections"] for y in s["years"] for b in y["blocks"])


# --------------------------------------------------------------------------- PDF
def build_pdf(model):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Frame, PageBreak,
                                    PageTemplate, Paragraph, Spacer, Table, TableStyle)

    page_w, page_h = A4
    info = model["header"]

    class NumberedCanvas(rl_canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved = []

        def showPage(self):
            self._saved.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved)
            for state in self._saved:
                self.__dict__.update(state)
                self._decorate(total)
                super().showPage()
            super().save()

        def _decorate(self, total):
            self.saveState()
            self.setLineWidth(0.8)
            self.rect(16, 16, page_w - 32.7, page_h - 32.1)
            if os.path.exists(LOGO_PATH):
                self.drawImage(LOGO_PATH, 29, page_h - 58, width=88, height=27, mask="auto")
            # Academic year (top right)
            tail = f"{info['academic_year']}  {info['semester']}"
            tail_w = stringWidth(tail, "Times-Roman", 13)
            self.setFont("Times-Roman", 13)
            self.drawRightString(page_w - 29, page_h - 34, tail)
            self.setFont("Times-Bold", 13)
            self.drawRightString(page_w - 29 - tail_w, page_h - 34, "Academic Year:")
            # Institute name (letter-spaced as in the reference)
            text = INSTITUTE_NAME
            natural = stringWidth(text, "Helvetica", 12.5)
            self.setFont("Helvetica", 12.5)
            self.drawString(127, page_h - 46, text, charSpace=(225 - natural) / (len(text) - 1))
            self.setFont("Courier-Bold", 8)
            self.drawCentredString(page_w / 2, page_h - 76, "=" * 88)
            # Title with the underlined phrase
            pre, mid = info["timetable_status"] + ": -", info["exam_title"]
            post = ""
            font, size = "Times-Bold", 16.5
            w_pre, w_mid, w_post = (stringWidth(t, font, size) for t in (pre, mid, post))
            x = (page_w - (w_pre + w_mid + w_post)) / 2
            y = page_h - 98
            self.setFont(font, size)
            self.drawString(x, y, pre)
            self.drawString(x + w_pre, y, mid)
            self.setLineWidth(0.8)
            self.line(x + w_pre, y - 2, x + w_pre + w_mid, y - 2)
            self.drawString(x + w_pre + w_mid, y, post)
            self.setFont("Times-Bold", info["summary_size"])
            if len(info["summary_lines"]) > 1:  # long summary: two lines, still clear of the table frame
                for offset, line in zip((107, 115), info["summary_lines"]):
                    self.drawCentredString(page_w / 2, page_h - offset, line)
            elif info["summary_lines"]:
                self.drawCentredString(page_w / 2, page_h - 112, info["summary_lines"][0])
            # Footer
            self.setFont("Times-Bold", 8)
            self.drawString(46, page_h - 799, f"Prepared by: {PREPARED_BY}")
            self.drawString(160, page_h - 799, f"Checked by: {CHECKED_BY}")
            self.drawString(290, page_h - 799, "Dy. Controller of Examinations")
            self.drawString(464, page_h - 799, "Controller of Examinations")
            self.drawString(46, page_h - 811, f"Date: {info['prepared_date']}")
            self.setFont("Helvetica", 7)
            self.drawRightString(538, page_h - 811, f"Page {self._pageNumber} of {total}")
            self.restoreState()

    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.6, leading=10.2, alignment=TA_LEFT)
    cell_c = ParagraphStyle("cellc", parent=cell, alignment=TA_CENTER)
    head = ParagraphStyle("head", fontName="Times-Bold", fontSize=9, leading=10.5, alignment=TA_CENTER)
    group = ParagraphStyle("group", fontName="Times-Bold", fontSize=10, leading=12, alignment=TA_LEFT)

    def para(text, style):
        return Paragraph(escape(text).replace("\n", "<br/>"), style)

    def group_table(year):
        data = [[para(year["title"], group)] + [""] * 5, [para(h, head) for h in HEADERS]]
        heights = [18, 15]
        styles = [
            ("GRID", (0, 0), (-1, -1), 0.7, colors.black),
            ("SPAN", (0, 0), (-1, 0)),
            ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.88, 0.88, 0.88)),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            ("TOPPADDING", (0, 0), (-1, 1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, 1), 2.5),
            ("LEFTPADDING", (3, 2), (3, -1), 2), ("RIGHTPADDING", (3, 2), (3, -1), 2),
        ]
        for block in year["blocks"]:
            start = len(data)
            for i, (branch, sem, course, code) in enumerate(block["rows"]):
                data.append([para(block["day_date"], cell_c) if i == 0 else "",
                             para(block["time"], cell_c) if i == 0 else "",
                             para(branch, cell), para(sem, cell_c), para(course, cell), para(code, cell_c)])
                heights.append(None)
            end = len(data) - 1
            if end > start:
                styles += [("SPAN", (0, start), (0, end)), ("SPAN", (1, start), (1, end))]
            data.append([""] * 6)  # thin separator row between sessions, as in the reference
            heights.append(7)
        table = Table(data, colWidths=COL_WIDTHS, rowHeights=heights, repeatRows=2, hAlign="LEFT")
        table.setStyle(TableStyle(styles))
        return table

    output = BytesIO()
    doc = BaseDocTemplate(output, pagesize=A4, title=info["timetable_status"], author=INSTITUTE_NAME)
    frame = Frame(29, 57, sum(COL_WIDTHS), page_h - 120 - 57, leftPadding=0, rightPadding=0,
                  topPadding=0, bottomPadding=0, id="body")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame])])
    story = []
    for s_index, section in enumerate(model["sections"]):
        for y_index, year in enumerate(section["years"]):
            if s_index or y_index:
                story.append(PageBreak())  # every programme / year section starts on its own page
            story.append(group_table(year))
    if not story:
        story.append(Paragraph("No timetable entries.", cell))
    doc.build(story, canvasmaker=NumberedCanvas)
    output.seek(0)
    return output


# --------------------------------------------------------------------------- Word
def build_docx(model):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    info = model["header"]
    document = Document()
    document.core_properties.title = info["timetable_status"]
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Pt(595.3), Pt(841.9)
    section.left_margin = section.right_margin = Pt(29)
    section.top_margin, section.bottom_margin = Pt(121), Pt(58)
    section.header_distance, section.footer_distance = Pt(22), Pt(22)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(8.6)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)

    # Page border like the reference.
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "16"), ("color", "000000")):
            el.set(qn(f"w:{k}"), v)
        borders.append(el)
    section._sectPr.find(qn("w:pgMar")).addnext(borders)

    def style_run(run, name="Arial", size=8.6, bold=False, underline=False):
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        run.font.size = Pt(size)
        run.bold = bold
        run.underline = underline
        run.font.color.rgb = RGBColor(0, 0, 0)

    def set_cell_width(cell, width):
        cell.width = Pt(width)

    def set_table_borders(table, sz="6", none=False):
        tbl_pr = table._tbl.tblPr
        el = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "nil" if none else "single")
            if not none:
                e.set(qn("w:sz"), sz)
                e.set(qn("w:space"), "0")
                e.set(qn("w:color"), "000000")
            el.append(e)
        tbl_pr.append(el)

    def fixed_layout(table, widths):
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        table._tbl.tblPr.append(layout)
        grid = table._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 20)))

    def add_field(paragraph, instruction, **font):
        for kind in ("begin", None, "separate", "text", "end"):
            run = paragraph.add_run()
            style_run(run, **font)
            if kind == "text":
                run.text = "1"
            elif kind is None:
                instr = OxmlElement("w:instrText")
                instr.set(qn("xml:space"), "preserve")
                instr.text = f" {instruction} "
                run._r.append(instr)
            else:
                fld = OxmlElement("w:fldChar")
                fld.set(qn("w:fldCharType"), kind)
                run._r.append(fld)

    # ---- header (repeats on every page, as in the reference)
    header = section.header
    header.is_linked_to_previous = False
    top = header.add_table(rows=1, cols=3, width=Pt(537.3))
    widths = [98, 227, 212.3]
    fixed_layout(top, widths)
    set_table_borders(top, none=True)
    c0, c1, c2 = top.rows[0].cells
    for c, w in zip((c0, c1, c2), widths):
        set_cell_width(c, w)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if os.path.exists(LOGO_PATH):
        c0.paragraphs[0].add_run().add_picture(LOGO_PATH, width=Pt(88), height=Pt(27))
    p = c1.paragraphs[0]
    r = p.add_run(INSTITUTE_NAME)
    style_run(r, "Arial", 12.5)
    r.font.spacing = Pt(0.6)
    p = c2.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(p.add_run("Academic Year:"), "Times New Roman", 13, bold=True)
    style_run(p.add_run(f"{info['academic_year']}  {info['semester']}"), "Times New Roman", 13)
    # python-docx puts a leading empty paragraph in a header; remove it
    first = header.paragraphs[0]
    first._element.getparent().remove(first._element)

    def header_line(text_runs, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0):
        par = header.add_paragraph()
        par.alignment = align
        par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(before), Pt(after)
        for text, kwargs in text_runs:
            style_run(par.add_run(text), **kwargs)
        return par

    header_line([("=" * 88, dict(name="Courier New", size=8, bold=True))], before=4)
    pre, mid = info["timetable_status"] + ": -", info["exam_title"]
    post = ""
    tnr = dict(name="Times New Roman", size=16.5, bold=True)
    header_line([(pre, tnr), (mid, dict(tnr, underline=True)), (post, tnr)], before=2)
    summary_par = header_line([(info["summary_lines"][0] if info["summary_lines"] else "",
                                dict(name="Times New Roman", size=info["summary_size"], bold=True))], after=4)
    for extra_line in info["summary_lines"][1:]:
        summary_par.runs[-1].add_break()
        style_run(summary_par.add_run(extra_line), "Times New Roman", info["summary_size"], bold=True)

    # ---- footer
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.style = document.styles["Normal"]  # drop the Footer style's built-in centre/right tab stops
    for tab in (131, 261, 435):
        fp.paragraph_format.tab_stops.add_tab_stop(Pt(tab), WD_TAB_ALIGNMENT.LEFT)
    fb = dict(name="Times New Roman", size=8, bold=True)
    style_run(fp.add_run(f"Prepared by: {PREPARED_BY}\tChecked by: {CHECKED_BY}\t"
                         "Dy. Controller of Examinations\tController of Examinations"), **fb)
    fp2 = footer.add_paragraph()
    fp2.paragraph_format.tab_stops.add_tab_stop(Pt(509), WD_TAB_ALIGNMENT.RIGHT)
    style_run(fp2.add_run(f"Date: {info['prepared_date']}"), **fb)
    small = dict(name="Arial", size=7)
    style_run(fp2.add_run("\tPage "), **small)
    add_field(fp2, "PAGE", **small)
    style_run(fp2.add_run(" of "), **small)
    add_field(fp2, "NUMPAGES", **small)

    # ---- body
    def write_cell(cell, text, width, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, font="Arial",
                   size=8.6, keep_next=False):
        set_cell_width(cell, width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        lines = str(text).split("\n") if text else [""]
        paragraph = cell.paragraphs[0]
        for i, line in enumerate(lines):
            if i:
                paragraph = cell.add_paragraph()
            paragraph.alignment = align
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = keep_next
            style_run(paragraph.add_run(line), font, size, bold)

    def shade(cell, fill):
        el = OxmlElement("w:shd")
        el.set(qn("w:val"), "clear")
        el.set(qn("w:color"), "auto")
        el.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(el)

    def row_props(row, header=False, height=None):
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        if height:
            h = OxmlElement("w:trHeight")
            h.set(qn("w:val"), str(int(height * 20)))
            h.set(qn("w:hRule"), "exact")
            trpr.append(h)
        if header:
            trpr.append(OxmlElement("w:tblHeader"))

    def cell_margins(table, top=50, bottom=50, left=100, right=100):
        mar = OxmlElement("w:tblCellMar")
        for name, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
            e = OxmlElement(f"w:{name}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        table._tbl.tblPr.append(mar)

    C, L = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    first_table = True
    for s_index, sec in enumerate(model["sections"]):
        for y_index, year in enumerate(sec["years"]):
            if first_table:
                first_table = False
            else:
                gap = document.add_paragraph()
                gap.paragraph_format.space_after = Pt(4)
                gap.paragraph_format.page_break_before = True  # every programme / year section -> new page
                style_run(gap.add_run(""), size=4)
            n_rows = 2 + sum(len(b["rows"]) + 1 for b in year["blocks"])
            table = document.add_table(rows=n_rows, cols=6)
            fixed_layout(table, COL_WIDTHS)
            set_table_borders(table)
            cell_margins(table)
            # group heading row (merged, repeats on each page)
            hr = table.rows[0]
            merged = hr.cells[0].merge(hr.cells[5])
            write_cell(merged, year["title"], sum(COL_WIDTHS), L, True, "Times New Roman", 10, True)
            shade(merged, "E0E0E0")
            row_props(hr, header=True)
            # column headings
            for i, h in enumerate(HEADERS):
                write_cell(table.rows[1].cells[i], h, COL_WIDTHS[i], C, True, "Times New Roman", 9, True)
            row_props(table.rows[1], header=True)
            r = 2
            for block in year["blocks"]:
                start = r
                count = len(block["rows"])
                for i, (branch, sem, course, code) in enumerate(block["rows"]):
                    row = table.rows[r]
                    keep = i < count  # keep the block together with its separator row
                    if i == 0:
                        write_cell(row.cells[0], block["day_date"], COL_WIDTHS[0], C, keep_next=keep)
                        write_cell(row.cells[1], block["time"], COL_WIDTHS[1], C, keep_next=keep)
                    write_cell(row.cells[2], branch, COL_WIDTHS[2], L, keep_next=keep)
                    write_cell(row.cells[3], sem, COL_WIDTHS[3], C, keep_next=keep)
                    write_cell(row.cells[4], course, COL_WIDTHS[4], L, keep_next=keep)
                    write_cell(row.cells[5], code, COL_WIDTHS[5], C, keep_next=keep)
                    row_props(row)
                    r += 1
                if count > 1:  # vertical merge of Day & Date and Time, as in the reference
                    for col in (0, 1):
                        top_cell = table.cell(start, col)
                        merged = top_cell.merge(table.cell(start + count - 1, col))
                        # merge() concatenates empty paragraphs; keep only the written ones
                        for extra in merged.paragraphs:
                            if not extra.text.strip() and len(merged.paragraphs) > 1:
                                extra._element.getparent().remove(extra._element)
                        merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                sep = table.rows[r]  # thin separator row between sessions
                for i in range(6):
                    set_cell_width(sep.cells[i], COL_WIDTHS[i])
                    par = sep.cells[i].paragraphs[0]
                    style_run(par.add_run(""), size=2)
                    par.paragraph_format.line_spacing = Pt(2)
                row_props(sep, height=7)
                r += 1
    if first_table:
        document.add_paragraph("No timetable entries.")
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output
