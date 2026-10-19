"""Flask application for the Re-Exam Scheduler."""
import json
import logging
import os
import secrets
import shutil
from datetime import date, datetime
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

import scheduler
from auth import init_auth, login_required

load_dotenv()
BASE_DIR = Path(__file__).resolve().parent
INSTANCE_DIR = Path(os.environ.get("REEXAM_INSTANCE_DIR", BASE_DIR / "instance")).resolve()
INSTANCE_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR = BASE_DIR / "generated_schedules"
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
META_FILE = OUTPUT_DIR / "timetable_meta.json"
TIMETABLE_FILE = OUTPUT_DIR / "generated_reexam_timetable.xlsx"
CLEANED_FILE = OUTPUT_DIR / "cleaned_reexam_database.xlsx"


def _secret_key():
    value = os.environ.get("SECRET_KEY", "").strip()
    if value:
        if len(value) < 32:
            raise RuntimeError("SECRET_KEY must be at least 32 characters long.")
        return value
    path = INSTANCE_DIR / "secret_key.txt"
    if path.exists():
        value = path.read_text(encoding="utf-8").strip()
        if len(value) >= 32:
            return value
    value = secrets.token_urlsafe(48)
    path.write_text(value, encoding="utf-8")
    try:
        path.chmod(0o600)
    except OSError:
        pass
    return value


app = Flask(__name__, template_folder=str(BASE_DIR / "templates"), static_folder=str(BASE_DIR / "static"))
app.config.update(
    SECRET_KEY=_secret_key(), MAX_CONTENT_LENGTH=30 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=1800,
)
if os.environ.get("COOKIE_SECURE", "").lower() in {"1", "true", "yes", "on"}:
    app.config["SESSION_COOKIE_SECURE"] = True
init_auth(app, str(INSTANCE_DIR))
logging.basicConfig(level=os.environ.get("LOG_LEVEL", "INFO"))


def _save_meta(meta):
    temp = META_FILE.with_suffix(".tmp")
    temp.write_text(json.dumps(meta, indent=2), encoding="utf-8")
    os.replace(temp, META_FILE)


def _load_timetable_meta():
    try:
        data = json.loads(META_FILE.read_text(encoding="utf-8"))
        return (data.get("academic_year", ""), data.get("semester", ""),
                data.get("exam_type", "Re-exam"), data.get("timetable_status", "draft"))
    except (OSError, ValueError, TypeError):
        return "", "", "Re-exam", "draft"


def _read_timetable_sheets(file_path):
    wb = load_workbook(file_path, data_only=True, read_only=True)
    try:
        result = []
        for ws in wb.worksheets:
            rows = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
            while rows and not any(str(v).strip() for v in rows[-1] if v is not None):
                rows.pop()
            if rows:
                result.append((ws.title, rows))
        return result
    finally:
        wb.close()


def _positive_capacity(value, label, default):
    if value is None or str(value).strip() == "":
        return default
    value = str(value).strip()
